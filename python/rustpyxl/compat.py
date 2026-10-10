"""Drop-in openpyxl API bridge: native supported reads, reference fallbacks.

This is a hybrid migration layer, not a complete Rust-only implementation.
Use `import rustpyxl.compat as openpyxl`, or call install() before existing imports.
Native interception is validated against openpyxl 3.1.5; other versions delegate.
"""
from os import fspath
from operator import index as integer_index
from types import MethodType
from weakref import WeakSet

try:
    import openpyxl as _reference
except ImportError as exc:
    raise ImportError("Install rustpyxl[compat] for drop-in openpyxl compatibility") from exc

from openpyxl.reader import excel as _excel
from openpyxl.xml.constants import SHARED_STRINGS
from ._native import NativeWorkbook

_original_load = _reference.load_workbook
_original_excel_load = _excel.load_workbook
Workbook = _reference.Workbook
__version__ = _reference.__version__


def __getattr__(name):
    return getattr(_reference, name)


class _SharedStrings:
    """Reference parsers can index Rust storage without a second Python table."""
    def __init__(self, native):
        self._native = native

    def __len__(self):
        return self._native.shared_len()

    def __getitem__(self, index):
        if isinstance(index, slice):
            return [self[i] for i in range(*index.indices(len(self)))]
        index = integer_index(index)
        if index < 0:
            index += len(self)
        if index < 0 or index >= len(self):
            raise IndexError("list index out of range")
        return self._native.shared_string(index)


class _State:
    def __init__(self, native, data_only):
        self.native = native
        self.data_only = data_only
        self.closed = False
        self.iterators = WeakSet()
        self.sheet_indices = {name: i for i, name in enumerate(native.sheetnames)}

    def close(self):
        self.closed = True
        for iterator in self.iterators:
            iterator.close()
        self.iterators.clear()
        self.native = None


def _accelerate_sheet(sheet, state):
    original = sheet._cells_by_row
    sheet._rustpyxl_accelerated = True

    def cells_by_row(self, min_col, min_row, max_col, max_row, values_only=False):
        width = max_col or self.max_column
        limit = max_row or self.max_row
        # Preserve actual ReadOnlyCell objects and openpyxl's dimension-free
        # gap behavior by using the reference parser for these operations.
        if not values_only or width is None:
            yield from original(min_col, min_row, max_col, max_row, values_only)
            return
        if state.closed:
            raise ValueError("workbook is closed")
        try:
            native = state.native.rows(state.sheet_indices[self.title], state.data_only)
        except (ValueError, KeyError):
            yield from original(min_col, min_row, max_col, max_row, values_only)
            return
        state.iterators.add(native)
        number = 0
        try:
            while limit is None or number < limit:
                if state.closed:
                    raise ValueError("workbook is closed")
                size = 1024 if limit is None else min(1024, limit - number)
                try:
                    batch = native.batch(size)
                except ValueError:
                    # A failed native batch yields no partial rows. Resume the
                    # original parser at the first row not yet delivered.
                    yield from original(min_col, max(min_row, number + 1), max_col, max_row, values_only)
                    return
                if not batch:
                    return
                for row in batch:
                    if state.closed:
                        raise ValueError("workbook is closed")
                    number += 1
                    if number < min_row:
                        continue
                    if len(row) < width:
                        row += (None,) * (width - len(row))
                    yield row[min_col - 1:width]
        finally:
            native.close()
            state.iterators.discard(native)

    sheet._cells_by_row = MethodType(cells_by_row, sheet)


def load_workbook(filename, read_only=False, keep_vba=False, data_only=False,
                  keep_links=True, rich_text=False):
    """Match openpyxl 3.1.5's public signature and return its workbook classes."""
    options = dict(read_only=read_only, keep_vba=keep_vba, data_only=data_only,
                   keep_links=keep_links, rich_text=rich_text)
    if not read_only or rich_text or _reference.__version__ != "3.1.5":
        return _original_load(filename, **options)
    try:
        path = fspath(filename)
        if not isinstance(path, str):
            return _original_load(filename, **options)
        native = NativeWorkbook(path)
    except (TypeError, ValueError, OSError):
        return _original_load(filename, **options)

    # Use the original loader's complete OOXML object construction. Substitute
    # only shared-string storage, so styles, dates, names, links, etc. retain
    # their reference behavior without duplicating the resident string table.
    reader = _excel.ExcelReader(filename, read_only, keep_vba, data_only, keep_links, rich_text)
    original_strings = reader.read_strings
    holder = [native]

    def read_strings(self):
        part = self.package.find(SHARED_STRINGS)
        manifest_path = part.PartName.lstrip("/") if part is not None else None
        if manifest_path != native.shared_strings_path:
            holder[0] = None
            original_strings()
        else:
            self.shared_strings = _SharedStrings(native)

    reader.read_strings = MethodType(read_strings, reader)
    try:
        reader.read()
    except Exception:
        reader.archive.close()
        raise
    workbook = reader.wb
    if holder[0] is None:
        return workbook
    state = _State(native, data_only)
    for sheet in workbook.worksheets:
        if sheet.title in state.sheet_indices:
            _accelerate_sheet(sheet, state)
    original_close = workbook.close

    def close(self):
        state.close()
        original_close()

    workbook.close = MethodType(close, workbook)
    workbook._rustpyxl_state = state
    return workbook


def install():
    """Opt in to acceleration for existing openpyxl imports in this process.

    Call at startup, before importing load_workbook by name in other modules.
    No installed files or global Python environments are modified.
    """
    _reference.load_workbook = load_workbook
    _excel.load_workbook = load_workbook


def uninstall():
    """Restore the original loaders; already loaded workbooks remain usable."""
    if _reference.load_workbook is load_workbook:
        _reference.load_workbook = _original_load
    if _excel.load_workbook is load_workbook:
        _excel.load_workbook = _original_excel_load
