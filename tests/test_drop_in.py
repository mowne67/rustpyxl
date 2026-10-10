import datetime
import io
import inspect
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import pytest
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from rustpyxl import compat
from test_compatibility import package, xml, reference_rows


def test_same_public_classes_and_signature():
    assert compat.Workbook is openpyxl.Workbook
    assert inspect.signature(compat.load_workbook) == inspect.signature(openpyxl.load_workbook)
    assert isinstance(compat.Workbook(), openpyxl.Workbook)


def test_mutable_workbook_styles_names_merges_roundtrip(tmp_path):
    path = tmp_path / "rich.xlsx"
    wb = compat.Workbook()
    ws = wb.active
    ws.title = "Styled"
    ws["A1"] = "heading"
    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["A1"].fill = PatternFill("solid", fgColor="00FF00")
    ws.merge_cells("A1:C1")
    ws["A2"] = datetime.datetime(2025, 2, 3, 4, 5, 6)
    ws["B2"] = "=SUM(1,2)"
    ws.freeze_panes = "A2"
    wb.defined_names.add(DefinedName("Heading", attr_text="'Styled'!$A$1"))
    wb.save(path)
    loaded = compat.load_workbook(path)
    assert loaded["Styled"]["A1"].font.bold
    assert str(loaded["Styled"].merged_cells) == "A1:C1"
    assert loaded["Styled"].freeze_panes == "A2"
    assert "Heading" in loaded.defined_names
    loaded["Styled"]["B3"] = 99
    loaded.save(tmp_path / "edited.xlsx")
    assert openpyxl.load_workbook(tmp_path / "edited.xlsx")["Styled"]["B3"].value == 99


def test_native_read_with_reference_cells_and_ranges(tmp_path):
    path = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 2050):
        ws.append([i, True, f"value-{i}", i * 0.5])
    wb.save(path)
    rb = compat.load_workbook(path, read_only=True)
    try:
        assert rb.active._rustpyxl_accelerated
        assert list(rb.active.iter_rows(values_only=True)) == reference_rows(path)
        bounds = dict(min_row=1000, max_row=2048, min_col=2, max_col=6)
        assert list(rb.active.iter_rows(values_only=True, **bounds)) == reference_rows(path, **bounds)
        assert isinstance(rb.active["A1"], ReadOnlyCell)
        assert rb.active["A1"].value == 1
        with ThreadPoolExecutor(max_workers=3) as executor:
            assert list(executor.map(lambda _: sum(r[0] for r in rb.active.values), range(3))) == [sum(range(1, 2050))] * 3
    finally:
        rb.close()


@pytest.mark.parametrize("data_only", [False, True])
def test_fallback_after_complete_batches_no_duplicates(tmp_path, data_only):
    path = tmp_path / "late-date.xlsx"
    wb = openpyxl.Workbook()
    for i in range(2200):
        wb.active.append([i, f"row-{i}"])
    wb.active["B2050"] = datetime.datetime(2025, 1, 2)
    wb.save(path)
    rb = compat.load_workbook(path, read_only=True, data_only=data_only)
    try:
        rows = list(rb.active.values)
        assert rows == reference_rows(path, data_only=data_only)
        assert isinstance(rows[2049][1], datetime.datetime)
    finally:
        rb.close()


def test_shared_formulas_and_dates_reference_fallback(tmp_path):
    content = '<row r="1"><c r="A1"><f t="shared" si="0" ref="A1:A2">B1+1</f><v>7</v></c></row><row r="2"><c r="A2"><f t="shared" si="0"/><v>8</v></c></row>'
    path = package(tmp_path, xml(content, '<dimension ref="A1:A2"/>'))
    for data_only in [True, False]:
        wb = compat.load_workbook(path, read_only=True, data_only=data_only)
        try:
            assert list(wb.active.values) == reference_rows(path, data_only=data_only)
        finally:
            wb.close()


def test_large_integer_falls_back_without_precision_loss(tmp_path):
    path = package(tmp_path, xml('<row r="1"><c r="A1"><v>9223372036854775809</v></c></row>', '<dimension ref="A1:A1"/>'))
    wb = compat.load_workbook(path, read_only=True)
    try:
        value = next(wb.active.values)[0]
        assert type(value) is int
        assert value == 9223372036854775809
    finally:
        wb.close()


def test_shared_strings_proxy_supports_native_and_reference_cells(tmp_path):
    rels = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet1.xml" Id="rId1"/><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="sst"/></Relationships>'
    sst = '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t> shared &amp; text </t></si></sst>'
    path = package(tmp_path, xml('<row r="1"><c r="A1" t="s"><v>0</v></c></row>', '<dimension ref="A1:A1"/>'), {"xl/sharedStrings.xml": sst, "xl/_rels/workbook.xml.rels": rels})
    rb = compat.load_workbook(path, read_only=True)
    try:
        assert not isinstance(rb.active._shared_strings, list)
        assert list(rb.active.values) == [(" shared & text ",)]
        assert rb.active["A1"].value == " shared & text "
        assert rb.active._shared_strings[-1] == " shared & text "
    finally:
        rb.close()


def test_file_like_and_rich_text_paths_delegate(tmp_path):
    path = tmp_path / "source.xlsx"
    wb = compat.Workbook()
    wb.active.append([1, "text"])
    wb.save(path)
    for options in [{"read_only": True}, {"read_only": True, "rich_text": True}]:
        rb = compat.load_workbook(io.BytesIO(path.read_bytes()), **options)
        assert list(rb.active.values) == [(1, "text")]
        assert not hasattr(rb, "_rustpyxl_state")
        rb.close()


def test_install_and_uninstall_are_process_local(tmp_path):
    original = openpyxl.load_workbook
    path = tmp_path / "install.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append([1])
    wb.save(path)
    try:
        compat.install()
        compat.install()
        rb = openpyxl.load_workbook(path, read_only=True)
        assert rb.active._rustpyxl_accelerated
        assert list(rb.active.values) == [(1,)]
        rb.close()
    finally:
        compat.uninstall()
    assert openpyxl.load_workbook is original


def test_close_releases_active_native_iterators(tmp_path):
    path = tmp_path / "close.xlsx"
    wb = openpyxl.Workbook()
    for i in range(1100):
        wb.active.append([i])
    wb.save(path)
    rb = compat.load_workbook(path, read_only=True)
    iterator = rb.active.values
    assert next(iterator) == (0,)
    rb.close()
    assert rb._rustpyxl_state.native is None
    with pytest.raises(ValueError, match="closed"):
        next(iterator)


def test_reset_dimensions_and_version_guard(tmp_path, monkeypatch):
    path = tmp_path / "dimensions.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append([1, 2])
    wb.save(path)
    rb = compat.load_workbook(path, read_only=True)
    rb.active.reset_dimensions()
    assert list(rb.active.values) == [(1, 2)]
    rb.close()
    monkeypatch.setattr(openpyxl, "__version__", "future-version")
    rb = compat.load_workbook(path, read_only=True)
    assert not hasattr(rb, "_rustpyxl_state")
    rb.close()
