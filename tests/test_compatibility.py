"""Differential integration tests against openpyxl, not implementation mocks."""
import datetime
import random
import zipfile
from pathlib import Path

import openpyxl
import pytest
import rustpyxl


def reference_rows(path, sheet=0, data_only=False, **bounds):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=data_only)
    try:
        # openpyxl returns [] for an empty gap without a declared dimension;
        # normalize that inconsistency to the tuple row interface.
        return [tuple(row) for row in wb.worksheets[sheet].iter_rows(values_only=True, **bounds)]
    finally:
        wb.close()


def native_rows(path, sheet=0, data_only=False, **bounds):
    with rustpyxl.load_workbook(path, read_only=True, data_only=data_only) as wb:
        return list(wb.worksheets[sheet].iter_rows(values_only=True, **bounds))


def package(tmp_path, worksheet, extras=None, compression=zipfile.ZIP_DEFLATED):
    base = tmp_path / "base.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append([1])
    wb.save(base)
    path = tmp_path / "fixture.xlsx"
    with zipfile.ZipFile(base) as source, zipfile.ZipFile(path, "w", compression=compression) as dest:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = worksheet.encode()
            if extras and name in extras:
                data = extras[name].encode()
            if extras and "xl/sharedStrings.xml" in extras and name == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            dest.writestr(name, data)
        for name, value in (extras or {}).items():
            if name not in source.namelist():
                dest.writestr(name, value)
    return path


def xml(content, dimension=""):
    return f'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">{dimension}<sheetData>{content}</sheetData></worksheet>'


@pytest.mark.parametrize("data_only", [False, True])
@pytest.mark.parametrize("batch_size", [1, 7, 1024])
def test_openpyxl_reader_differential(tmp_path, data_only, batch_size):
    rng = random.Random(42)
    values = [None, True, False, 0, -42, 1234567890123, 0.25, -3.75, "hello", " <&>\"' ", "தமிழ் 🐍", "line\nbreak", "=A1+1", "#DIV/0!", "=", ""]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data & "quotes"'
    for i in range(1, 130):
        if i % 13:
            for j in range(1, 14):
                value = rng.choice(values)
                if value is not None:
                    ws.cell(i, j, value)
    ws.cell(133, 13, True)
    path = tmp_path / "source.xlsx"
    wb.save(path)
    assert native_rows(path, data_only=data_only, batch_size=batch_size) == reference_rows(path, data_only=data_only)
    bounds = dict(min_row=9, max_row=140, min_col=3, max_col=16)
    assert native_rows(path, data_only=data_only, batch_size=batch_size, **bounds) == reference_rows(path, data_only=data_only, **bounds)


def test_writer_roundtrip_and_openpyxl_interop(tmp_path):
    path = tmp_path / "written.xlsx"
    rows = [(1, True, " <&> ", None, "=A1+1"), (None, False, "தமிழ்", 2.5), ("=", "#N/A", "line\nbreak")]
    wb = rustpyxl.Workbook(write_only=True)
    ws = wb.create_sheet('Sales & "tax"')
    ws.append_rows(iter(rows), batch_size=2)
    second = wb.create_sheet("Second")
    second.append([9, "other"])
    wb.save(path)
    assert native_rows(path) == reference_rows(path)
    assert native_rows(path, 1) == [(9, "other")]
    with rustpyxl.load_workbook(path, read_only=True) as rb:
        assert rb.sheetnames == ['Sales & "tax"', "Second"]
    assert reference_rows(path)[0] == (1, True, " <&> ", None, "=A1+1")
    assert reference_rows(path, data_only=True)[0][-1] is None
    with pytest.raises(ValueError):
        ws.append([1])
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None


def test_shared_strings_rich_text_and_cached_formula(tmp_path):
    rels = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet1.xml" Id="rId1"/><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="sst"/></Relationships>'
    sst = '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t xml:space="preserve"> A&amp;B _x005F_x000D_ </t></si><si><r><t>rich</t></r><r><t>&#32;text</t></r><rPh sb="0" eb="1"><t>phonetic</t></rPh></si></sst>'
    path = package(tmp_path, xml('<row r="2"><c r="B2" t="s"><v>0</v></c><c r="D2" t="s"><v>1</v></c><c r="E2"><f>1+2</f><v>3</v></c><c r="F2" t="str"><f>"x"</f><v>cached</v></c></row>', '<dimension ref="A1:F4"/>'), {"xl/sharedStrings.xml": sst, "xl/_rels/workbook.xml.rels": rels})
    for data_only in [True, False]:
        assert native_rows(path, data_only=data_only) == reference_rows(path, data_only=data_only)


@pytest.mark.parametrize("compression", [zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED])
def test_sparse_no_refs_and_cdata(tmp_path, compression):
    path = package(tmp_path, xml('<row r="3"><c t="inlineStr"><is><t><![CDATA[ <&> ]]></t></is></c><c><v>9</v></c></row><row r="5"><c r="C5" t="b"><v>0</v></c></row>'), compression=compression)
    assert native_rows(path) == reference_rows(path)


def test_unsupported_dates_fail_explicitly(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append([datetime.datetime(2025, 1, 2)])
    path = tmp_path / "dates.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="style engine"):
        native_rows(path)


def test_shared_formula_rejected_but_cached_values_supported(tmp_path):
    path = package(tmp_path, xml('<row r="1"><c r="A1"><f t="shared" si="0" ref="A1:A2">B1+1</f><v>7</v></c></row><row r="2"><c r="A2"><f t="shared" si="0"/><v>8</v></c></row>'))
    with pytest.raises(ValueError, match="shared formula"):
        native_rows(path)
    assert native_rows(path, data_only=True) == [(7,), (8,)]


def test_bad_shared_string_index(tmp_path):
    path = package(tmp_path, xml('<row r="1"><c r="A1" t="s"><v>20</v></c></row>'))
    with pytest.raises(ValueError, match="shared string index"):
        native_rows(path)


def test_checksum_corruption_is_detected(tmp_path):
    path = package(tmp_path, xml('<row r="1"><c r="A1"><v>123</v></c></row>'), compression=zipfile.ZIP_STORED)
    data = path.read_bytes()
    assert b"<v>123</v>" in data
    path.write_bytes(data.replace(b"<v>123</v>", b"<v>124</v>"))
    with pytest.raises(ValueError, match="CRC"):
        native_rows(path)


def test_truncated_xml(tmp_path):
    path = package(tmp_path, '<worksheet><sheetData><row r="1"><c r="A1"><v>9</v>')
    with pytest.raises(ValueError):
        native_rows(path)


@pytest.mark.parametrize("batch_size", [1, 1024])
def test_independent_iterators_and_close(tmp_path, batch_size):
    path = tmp_path / "multi.xlsx"
    wb = rustpyxl.Workbook(write_only=True)
    wb.create_sheet().append_rows([(i,) for i in range(20)])
    wb.save(path)
    rb = rustpyxl.load_workbook(path, read_only=True)
    a = rb.active.iter_rows(values_only=True, batch_size=batch_size)
    b = rb.active.iter_rows(values_only=True, batch_size=2)
    assert next(a) == next(b) == (0,)
    assert next(a) == next(b) == (1,)
    rb.close()
    with pytest.raises(ValueError, match="closed"):
        next(a)


def test_save_retry_and_no_partial_destination(tmp_path):
    path = tmp_path / "safe.xlsx"
    wb = rustpyxl.Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append([1])
    with pytest.raises(ValueError):
        wb.save(tmp_path / "missing" / "failed.xlsx")
    ws.append([2])
    wb.save(path)
    assert reference_rows(path) == [(1,), (2,)]


@pytest.mark.parametrize("name", ["", "x" * 32, "bad/name", "bad\x01name"])
def test_invalid_sheet_names(name):
    wb = rustpyxl.Workbook(write_only=True)
    with pytest.raises(ValueError):
        wb.create_sheet(name)


@pytest.mark.parametrize("value", [float("nan"), float("inf"), "bad\x01text", "x" * 32768, "=" + "x" * 32767, datetime.date(2025, 1, 1)])
def test_invalid_and_unsupported_values(tmp_path, value):
    wb = rustpyxl.Workbook(write_only=True)
    ws = wb.create_sheet()
    with pytest.raises((ValueError, TypeError)):
        ws.append([value])
    ws.append(["valid"])
    path = tmp_path / "valid.xlsx"
    wb.save(path)
    assert reference_rows(path) == [("valid",)]


def test_explicit_api_subset():
    with pytest.raises(NotImplementedError):
        rustpyxl.Workbook()
    with pytest.raises(NotImplementedError):
        rustpyxl.load_workbook("unused.xlsx")
    wb = rustpyxl.Workbook(write_only=True)
    wb.create_sheet("One")
    with pytest.raises(ValueError, match="duplicate"):
        wb.create_sheet("one")


def test_value_types_and_empty_cached_string(tmp_path):
    path = package(tmp_path, xml('<row r="1"><c r="A1" t="b"><v>1</v></c><c r="B1"><v>1</v></c><c r="C1"><v>1.5</v></c><c r="D1" t="str"><f>""</f><v></v></c></row>'))
    actual = native_rows(path, data_only=True)
    assert actual == reference_rows(path, data_only=True)
    assert [type(v) for v in actual[0]] == [bool, int, float, type(None)]


def test_carriage_return_and_empty_writer_values(tmp_path):
    wb = rustpyxl.Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["line\r\nnext", "", None, "end"])
    path = tmp_path / "newlines.xlsx"
    wb.save(path)
    assert native_rows(path) == reference_rows(path) == [("line\r\nnext", None, None, "end")]


@pytest.mark.parametrize("row", ["bad", b"bad", {"A": 1}])
def test_non_row_iterables_rejected(row):
    wb = rustpyxl.Workbook(write_only=True)
    with pytest.raises(TypeError):
        wb.create_sheet().append(row)


def test_sheet_name_whitespace_roundtrip(tmp_path):
    name = "Tab\tNew\nReturn\r"
    wb = rustpyxl.Workbook(write_only=True)
    wb.create_sheet(name).append([1])
    path = tmp_path / "title.xlsx"
    wb.save(path)
    with rustpyxl.load_workbook(path, read_only=True) as source:
        assert source.sheetnames == [name]
    reference = openpyxl.load_workbook(path, read_only=True)
    try:
        assert reference.sheetnames == [name]
    finally:
        reference.close()


def test_large_integer_read_is_explicitly_rejected(tmp_path):
    path = package(tmp_path, xml('<row r="1"><c r="A1"><v>9223372036854775809</v></c></row>'))
    with pytest.raises(ValueError, match="signed 64-bit"):
        native_rows(path)

